import openpyxl
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.table import Table, TableStyleInfo
from typing import List, Dict, Union, Self


class ExcelManager:

    def __init__(self) -> None:
        self.wb: Workbook = openpyxl.Workbook()
        self.sheet: Worksheet = self.wb.active
        self.header: List[Dict[str, Union[str, int]]] = []
        self.records: List[Dict[str, Union[str, int]]] = []
        self.starting_cell: str = ''
        self.ending_cell: str = ''
        self.range: Union[str, None] = None
        self.table: Union[Table, None] = None

    def read(self) -> None:
        pass

    def add_header(self, headers: List[Dict[str, str]]) -> Self:
        self.header = headers
        for header in headers:
            self.sheet.append([header['label'], header['value']])
        return self

    def add_blank(self) -> Self:
        self.sheet.append([])
        return self

    def add_data_records_to_datatable(self, records: List[Dict[str, Union[str, int]]]) -> Self:
        self.records = records
        self.sheet.append(list(records[0].keys()))
        for record in records:
            self.sheet.append(list(record.values()))
        return self

    def create_range(self) -> Self:
        start_cell_column: str = 'A'
        start_cell_row: int = len(self.header) + 2
        self.starting_cell = f"{start_cell_column}{str(start_cell_row)}"

        end_cell_column: str = chr(ord(start_cell_column) + len(list(self.records[0].keys())) - 1)
        end_cell_row: int = start_cell_row + len(self.records)
        self.ending_cell = f"{end_cell_column}{str(end_cell_row)}"

        self.range = f'{self.starting_cell}:{self.ending_cell}'

        return self

    def make_table(self, table_name: str) -> Self:
        self.table = Table(displayName=table_name, ref=self.range)
        return self

    def set_table_style(self) -> Self:
        self.table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=True)

        return self

    def add_table_to_sheet(self) -> Self:
        self.sheet.add_table(self.table)
        return self

    def save(self, file_name: str) -> Self:
        self.wb.save(file_name)
        return self


if __name__ == '__main__':
    excel_manager = ExcelManager()

    data_headers = [
        {'label': 'Title', 'value': ''},
        {'label': 'Pattern_A', 'value': 'A'},
        {'label': 'Pattern_B', 'value': 'B'},
        {'label': 'Pattern_C', 'value': 'C'},
        {'label': 'Pattern_D', 'value': '1'},
        {'label': 'Pattern_E', 'value': '2'},
        {'label': 'Pattern_F', 'value': '3'},
    ]

    data_records = [
        {'name': 'Bob', 'age': 20, 'Address': 'xxx'},
        {'name': 'Alice', 'age': 21, 'Address': 'xxx'},
        {'name': 'John', 'age': 22, 'Address': 'xxx'},
    ]

    (excel_manager
     .add_header(data_headers)
     .add_blank()
     .add_data_records_to_datatable(data_records)
     .create_range()
     .make_table('Table20')
     .set_table_style()
     .add_table_to_sheet().save('my_excel_.xlsx'))
